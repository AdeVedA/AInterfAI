import re
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

from docx import Document as _Docx
from ebooklib import ITEM_DOCUMENT, epub
from pdfminer.high_level import extract_text as _extract_pdf
from pptx import Presentation as _Pptx
from striprtf.striprtf import rtf_to_text as _extract_rtf

SUPPORTED = {".pdf", ".docx", ".pptx", ".rtf", ".txt", ".epub"}


def extract_text(path: Path) -> str:
    """useful function to extract textdata from different file formats through adapted modules"""
    suf = path.suffix.lower()
    if suf == ".pdf":
        print("pdf detected, processing...")
        pdf_file = open(f"{str(path)}", "rb")
        return _extract_pdf(pdf_file) or ""
    if suf == ".docx":
        print("docx detected, processing...")
        return "\n".join(p.text for p in _Docx(str(path)).paragraphs if p.text)
    if suf == ".pptx":
        print("pptx detected, processing...")
        prs = _Pptx(str(path))
        return "\n".join(
            shape.text
            for slide in prs.slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
    if suf == ".rtf":
        print("rtf detected, processing...")
        return _extract_rtf(path.read_text(errors="ignore"))
    if suf == ".epub":
        print("epub detected, processing...")

        class _EpubParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.out, self.skip = [], False

            def handle_starttag(self, tag, attrs):
                tag = tag.lower()
                if tag in {"sup", "script", "style"}:
                    self.skip = True
                if tag in {"p", "div", "br", "li"}:
                    self.out.append("\n")

            def handle_endtag(self, tag):
                if tag in {"sup", "script", "style"}:
                    self.skip = False
                if tag in {"p", "div", "li"}:
                    self.out.append("\n")

            def handle_data(self, data):
                if not self.skip:
                    self.out.append(data)

            def get_text(self):
                return "".join(self.out)

        book = epub.read_epub(str(path))
        texts = []
        for item in book.get_items_of_type(ITEM_DOCUMENT):
            raw = item.get_body_content().decode("utf-8", errors="ignore")
            parser = _EpubParser()
            parser.feed(raw)
            text = unescape(parser.get_text())
            # Nettoyage simple
            lines = [ln.strip() for ln in text.splitlines()]
            # on supprime les lignes vides et les numéros seuls (notes/pages)
            lines = [ln for ln in lines if ln and not re.fullmatch(r"\d{1,3}", ln)]
            if lines:
                texts.append("\n".join(lines))
        return "\n\n".join(texts)
    if suf in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    # et sinon :
    return ""
